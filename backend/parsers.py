"""
Excel report parsing - Posist "Payment Report" / "Bill Item Detailed Report" /
"Discount and Voucher Report" exports -> plain dict rows.

Copied from build_dashboard_data.py unchanged (same column layout, same
cleaning rules) so the API produces byte-identical rows to the old CLI
pipeline. If Posist ever changes a report's column layout, fix it here.
"""

import os
import glob
from datetime import datetime, date
from collections import defaultdict

import openpyxl


def find_branch_file(report_dir, branch_hint):
    """Find the xlsx for a branch inside a report folder, tolerant of typos
    in filenames (e.g. 'OTher Details chennai.xlsx')."""
    for fp in glob.glob(os.path.join(report_dir, "*.xlsx")):
        if branch_hint.lower() in os.path.basename(fp).lower():
            return fp
    raise FileNotFoundError(f"No file matching '{branch_hint}' in {report_dir}")


def find_dir_ci(parent, name):
    """Case-insensitive lookup of a subfolder named `name` inside `parent`.

    Real-world Posist/Google-Drive exports aren't always consistent about
    capitalization (e.g. a folder actually named "Discount and voucher
    Report" instead of "Discount and Voucher Report"), and unlike Windows,
    directory lookups are case-SENSITIVE on the Linux servers this runs on
    in production - a folder name that looks fine to a human eye can
    silently fail to be found, making real uploaded data show up as
    "not available" even though it's genuinely in the file. Falls back to
    the exact-case path if no case-insensitive match exists either (which
    then correctly reports as actually missing)."""
    exact = os.path.join(parent, name)
    if os.path.isdir(exact):
        return exact
    if not os.path.isdir(parent):
        return exact
    target = name.lower()
    for entry in os.listdir(parent):
        if entry.lower() == target and os.path.isdir(os.path.join(parent, entry)):
            return os.path.join(parent, entry)
    return exact


def _to_iso_date(value):
    """Payment/Bill-Item date-times come as 'DD-MM-YYYY hh:mm:ss pm' or
    'DD-Mon-YYYY hh:mm:ss pm' strings. Return 'YYYY-MM-DD'."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    for fmt in ("%d-%m-%Y %I:%M:%S %p", "%d-%b-%Y %I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _to_hour(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.hour
    s = str(value).strip()
    for fmt in ("%d-%m-%Y %I:%M:%S %p", "%d-%b-%Y %I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).hour
        except ValueError:
            continue
    return None


SERVICE_TYPE_BANQUET_NAMES = {"BANQUETS", "ODC"}


def _service_type(tab_name, tab_type):
    tab_name = (tab_name or "").strip()
    tab_type = (tab_type or "").strip().lower()
    if tab_name.upper() in SERVICE_TYPE_BANQUET_NAMES:
        return "Banquets & Catering"
    if tab_type == "delivery":
        return "Home Delivery"
    if tab_type == "takeout":
        return "Takeaway"
    return "Dine-In"


def parse_payment_report(fp, branch_label):
    """-> (allBills rows, {bill_no: total_amount}) for discount join."""
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    bills = []
    bill_totals = {}
    for row in ws.iter_rows(min_row=7, values_only=True):
        bill_no = row[1]
        if bill_no is None:
            continue  # date-section marker row
        void_flag = str(row[45]).strip().upper() if row[45] is not None else "NO"
        if void_flag == "YES":
            continue
        open_time = row[4]
        row_date = _to_iso_date(open_time)
        hour = _to_hour(open_time)
        covers = row[8] or 0
        net_sales = row[13] if row[13] is not None else 0.0
        total_amount = row[11] if row[11] is not None else 0.0
        tab_name = row[6]
        tab_type = row[7]
        bill_totals[str(bill_no).strip()] = float(total_amount)
        if row_date is None or hour is None:
            continue
        bills.append({
            "date": row_date,
            "hour": hour,
            "branch": branch_label,
            "bill": bill_no,
            "covers": int(covers) if isinstance(covers, (int, float)) else 0,
            "salesValue": float(net_sales),
            "serviceType": _service_type(tab_name, tab_type),
        })
    wb.close()
    return bills, bill_totals


def _clean_name(value):
    """Posist's own exports are inconsistent about internal whitespace in
    item/category names (e.g. 'Erachi      curry' vs 'Erachi curry'), which
    would otherwise split one dish into multiple chart rows."""
    return " ".join(str(value).split()) if value else value


def parse_bill_item_report(fp, branch_label):
    """-> aggregated salesData rows (per date+branch+category+item)."""
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    agg = defaultdict(lambda: {"qty": 0.0, "value": 0.0})
    for row in ws.iter_rows(min_row=9, values_only=True):
        qty = row[16]
        if not isinstance(qty, (int, float)) or qty <= 0:
            continue  # header repeats / BILL TOTAL rows / combo-constituent sub-lines
        category = _clean_name(row[10])
        item = _clean_name(row[11])
        amount = row[18] if isinstance(row[18], (int, float)) else 0.0
        row_date = _to_iso_date(row[7])  # Open Time
        if row_date is None or not item:
            continue
        key = (row_date, category, item)
        agg[key]["qty"] += qty
        agg[key]["value"] += amount
    wb.close()

    sales = []
    for (row_date, category, item), v in agg.items():
        rate = round(v["value"] / v["qty"], 2) if v["qty"] else 0.0
        sales.append({
            "date": row_date,
            "branch": branch_label,
            "category": category,
            "item": item,
            "qty": v["qty"] if v["qty"] % 1 else int(v["qty"]),
            "rate": rate,
            "value": round(v["value"], 2),
        })
    return sales


def _classify_discount(description, is_foc):
    desc = (description or "").strip()
    if desc == "-" or desc == "":
        return "Not Specified"
    if desc.lower().startswith("express offer(100%)"):
        return "Express Offer (100%)"
    if desc.lower().startswith("express"):
        return "Express Offer (Other)"
    if "zomato" in desc.lower():
        return "Zomato Merchant Discount"
    return "Not Specified"


def _discount_date(value):
    """The Discount and Voucher Report's date column is already plain text
    ('YYYY-MM-DD' or similar), unlike the datetime-stamped Payment/Bill-Item
    reports - so no _to_iso_date parsing here, just a datetime/date -> string
    safety net in case openpyxl hands back a real date object."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return value


def parse_discount_report(fp, branch_label, bill_totals):
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    discounts = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        trxno = row[4]
        if not trxno:
            continue
        row_date = _discount_date(row[3])
        disc_amt = row[7] if isinstance(row[7], (int, float)) else 0.0
        is_foc = str(row[11]).strip().upper() == "YES" if row[11] is not None else False
        item_value = bill_totals.get(str(trxno).strip())
        if item_value is None:
            # bill not found in this month's Payment Report (edge case); fall
            # back to treating the discount amount as the item value.
            item_value = disc_amt
        net = round(item_value - disc_amt, 2)
        discounts.append({
            "date": row_date,
            "branch": branch_label,
            "bill": trxno,
            "itemValue": round(item_value, 2),
            "discountValue": round(disc_amt, 2),
            "net": net,
            "isFull": is_foc,
            "reason": _classify_discount(row[6], is_foc),
            "type": row[6],
            "remarks": row[9],
        })
    wb.close()
    return discounts


BRANCH_MAP = {
    "Bangalore": "Bengaluru",
    "Chennai": "Chennai",
}


def parse_posist_export(posist_root):
    """Parse a whole Posist export folder -> (bills, sales, discounts, missing)
    combined across both branches.

    A missing report type or branch file (e.g. no "Discount and Voucher
    Report" this month, or one branch didn't send its Bill Item file) is
    tolerated: that section just comes back empty and its description is
    added to `missing`, instead of failing the whole upload - the dashboard
    then shows "no data" for whatever wasn't included rather than rejecting
    bills/sales that WERE there. Only raises if there isn't a single bill
    anywhere (Payment Report for both branches missing), since without any
    bills there's no date range to anchor the upload to."""
    payment_dir = find_dir_ci(posist_root, "Payment Report")
    bill_item_dir = find_dir_ci(posist_root, "Bill Item Detailed Report")
    discount_dir = find_dir_ci(posist_root, "Discount and Voucher Report")

    all_bills, sales_data, discount_data = [], [], []
    bill_totals_by_branch = {}
    missing = []

    for file_hint, branch_label in BRANCH_MAP.items():
        try:
            fp = find_branch_file(payment_dir, file_hint)
        except FileNotFoundError:
            missing.append(f"Payment Report ({branch_label})")
            bill_totals_by_branch[branch_label] = {}
            continue
        bills, bill_totals = parse_payment_report(fp, branch_label)
        all_bills.extend(bills)
        bill_totals_by_branch[branch_label] = bill_totals

    if not all_bills:
        raise FileNotFoundError(
            "No 'Payment Report' found for either branch - can't tell which dates this upload covers."
        )

    for file_hint, branch_label in BRANCH_MAP.items():
        try:
            fp = find_branch_file(bill_item_dir, file_hint)
        except FileNotFoundError:
            missing.append(f"Bill Item Detailed Report ({branch_label})")
            continue
        sales_data.extend(parse_bill_item_report(fp, branch_label))

    for file_hint, branch_label in BRANCH_MAP.items():
        try:
            fp = find_branch_file(discount_dir, file_hint)
        except FileNotFoundError:
            missing.append(f"Discount and Voucher Report ({branch_label})")
            continue
        discount_data.extend(parse_discount_report(fp, branch_label, bill_totals_by_branch[branch_label]))

    return all_bills, sales_data, discount_data, missing
