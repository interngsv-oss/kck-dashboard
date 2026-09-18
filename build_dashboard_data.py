#!/usr/bin/env python3
"""
KCK Sales Dashboard - data pipeline.

Reads the raw Posist export ("Posist Report" folder, unzipped from the
"Posist Report-*.zip" Google Drive export) and rebuilds the DASHBOARD_DATA
JSON blob embedded in kck_sales_dashboard_V11.html.

Usage:
    python build_dashboard_data.py "<path to Posist Report folder>" "<path to dashboard html>"

Example:
    python build_dashboard_data.py "C:\\Users\\Lenovo\\Downloads\\Posist Report" "C:\\Users\\Lenovo\\Downloads\\kck_sales_dashboard_V11.html"

What it rebuilds from the raw Excel reports:
    - allBills      <- Payment Report/{Bangalore,Chennai}.xlsx
    - salesData     <- Bill Item Detailed Report/{Bangalore,Chennai}.xlsx  (aggregated per day+item)
    - discountData  <- Discount and Voucher Report/{Bangalore,Chennai}.xlsx
                       joined against Payment Report for the pre-discount bill value

What it CANNOT rebuild yet:
    - cancelData (the "Cancelled KOTs" section). None of the 13 files in the
      Posist Report export contain item-level cancellation data (name/qty/rate
      of the cancelled dish). The KOT Tracking Report only has bill-level
      "KOT Voided" / "Bill Voided" status, not per-item detail. The old
      "Formatted.xlsx" mapping sheet expects a separate "KOT Report" export
      with columns: State, Date, Bill Number, Item Number, KOT Comment,
      Quantity, Rate, Amount - that report was not in this export.
      This script therefore carries over whatever cancelData already exists
      in the target HTML unchanged, and prints a warning. Export that report
      from Posist and this script can be extended to parse it.
"""

import sys
import os
import re
import json
import glob
from datetime import datetime
from collections import defaultdict

import openpyxl


# --------------------------------------------------------------------------
# Excel report parsing
# --------------------------------------------------------------------------

def _find_branch_file(report_dir, branch_hint):
    """Find the xlsx for a branch inside a report folder, tolerant of typos
    in filenames (e.g. 'OTher Details chennai.xlsx')."""
    for fp in glob.glob(os.path.join(report_dir, "*.xlsx")):
        if branch_hint.lower() in os.path.basename(fp).lower():
            return fp
    raise FileNotFoundError(f"No file matching '{branch_hint}' in {report_dir}")


def _to_iso_date(value):
    """Payment/Bill-Item date-times come as 'DD-MM-YYYY hh:mm:ss pm' or
    'DD-Mon-YYYY hh:mm:ss pm' strings. Return 'YYYY-MM-DD'."""
    if value is None:
        return None
    s = str(value).strip()
    for fmt in ("%d-%m-%Y %I:%M:%S %p", "%d-%b-%Y %I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _to_hour(value):
    if value is None:
        return None
    s = str(value).strip()
    for fmt in ("%d-%m-%Y %I:%M:%S %p", "%d-%b-%Y %I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).hour
        except ValueError:
            continue
    return None


SERVICE_TYPE_BANQUET_NAMES = {"BANQUETS", "ODC"}


def _service_type(tab_name, tab_type):
    tab_name = (tab_name or "").strip()
    tab_type = (tab_type or "").strip().lower()
    if tab_name.upper() in SERVICE_TYPE_BANQUET_NAMES:
        return "Banquets & Catering"
    if tab_type == "delivery":
        return "Home Delivery"
    if tab_type == "takeout":
        return "Takeaway"
    return "Dine-In"


def parse_payment_report(fp, branch_label):
    """-> (allBills rows, {bill_no: total_amount}) for discount join."""
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    bills = []
    bill_totals = {}
    for row in ws.iter_rows(min_row=7, values_only=True):
        bill_no = row[1]
        if bill_no is None:
            continue  # date-section marker row
        void_flag = str(row[45]).strip().upper() if row[45] is not None else "NO"
        if void_flag == "YES":
            continue
        open_time = row[4]
        date = _to_iso_date(open_time)
        hour = _to_hour(open_time)
        covers = row[8] or 0
        net_sales = row[13] if row[13] is not None else 0.0
        total_amount = row[11] if row[11] is not None else 0.0
        tab_name = row[6]
        tab_type = row[7]
        bill_totals[str(bill_no).strip()] = float(total_amount)
        if date is None or hour is None:
            continue
        bills.append({
            "date": date,
            "hour": hour,
            "branch": branch_label,
            "bill": bill_no,
            "covers": int(covers) if isinstance(covers, (int, float)) else 0,
            "salesValue": float(net_sales),
            "serviceType": _service_type(tab_name, tab_type),
        })
    wb.close()
    return bills, bill_totals


def _clean_name(value):
    """Posist's own exports are inconsistent about internal whitespace in
    item/category names (e.g. 'Erachi      curry' vs 'Erachi curry'), which
    would otherwise split one dish into multiple chart rows."""
    return " ".join(str(value).split()) if value else value


def parse_bill_item_report(fp, branch_label):
    """-> aggregated salesData rows (per date+branch+category+item)."""
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    agg = defaultdict(lambda: {"qty": 0.0, "value": 0.0})
    for row in ws.iter_rows(min_row=9, values_only=True):
        qty = row[16]
        if not isinstance(qty, (int, float)) or qty <= 0:
            continue  # header repeats / BILL TOTAL rows / combo-constituent sub-lines
        category = _clean_name(row[10])
        item = _clean_name(row[11])
        amount = row[18] if isinstance(row[18], (int, float)) else 0.0
        date = _to_iso_date(row[7])  # Open Time
        if date is None or not item:
            continue
        key = (date, category, item)
        agg[key]["qty"] += qty
        agg[key]["value"] += amount
    wb.close()

    sales = []
    for (date, category, item), v in agg.items():
        rate = round(v["value"] / v["qty"], 2) if v["qty"] else 0.0
        sales.append({
            "date": date,
            "branch": branch_label,
            "category": category,
            "item": item,
            "qty": v["qty"] if v["qty"] % 1 else int(v["qty"]),
            "rate": rate,
            "value": round(v["value"], 2),
        })
    return sales


def _classify_discount(description, is_foc):
    desc = (description or "").strip()
    if desc == "-" or desc == "":
        return "Not Specified"
    if desc.lower().startswith("express offer(100%)"):
        return "Express Offer (100%)"
    if desc.lower().startswith("express"):
        return "Express Offer (Other)"
    if "zomato" in desc.lower():
        return "Zomato Merchant Discount"
    return "Not Specified"


def parse_discount_report(fp, branch_label, bill_totals):
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    discounts = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        trxno = row[4]
        if not trxno:
            continue
        date = row[3]
        disc_amt = row[7] if isinstance(row[7], (int, float)) else 0.0
        is_foc = str(row[11]).strip().upper() == "YES" if row[11] is not None else False
        item_value = bill_totals.get(str(trxno).strip())
        if item_value is None:
            # bill not found in this month's Payment Report (edge case); fall
            # back to treating the discount amount as the item value.
            item_value = disc_amt
        net = round(item_value - disc_amt, 2)
        discounts.append({
            "date": date,
            "branch": branch_label,
            "bill": trxno,
            "itemValue": round(item_value, 2),
            "discountValue": round(disc_amt, 2),
            "net": net,
            "isFull": is_foc,
            "reason": _classify_discount(row[6], is_foc),
            "type": row[6],
            "remarks": row[9],
        })
    wb.close()
    return discounts


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

BRANCH_MAP = {
    "Bangalore": "Bengaluru",
    "Chennai": "Chennai",
}


def build_dataset(posist_root, existing_data):
    all_bills = []
    sales_data = []
    discount_data = []

    payment_dir = os.path.join(posist_root, "Payment Report")
    bill_item_dir = os.path.join(posist_root, "Bill Item Detailed Report")
    discount_dir = os.path.join(posist_root, "Discount and Voucher Report")

    bill_totals_by_branch = {}

    for file_hint, branch_label in BRANCH_MAP.items():
        fp = _find_branch_file(payment_dir, file_hint)
        bills, bill_totals = parse_payment_report(fp, branch_label)
        all_bills.extend(bills)
        bill_totals_by_branch[branch_label] = bill_totals
        print(f"  Payment Report [{branch_label}]: {len(bills)} bills")

    for file_hint, branch_label in BRANCH_MAP.items():
        fp = _find_branch_file(bill_item_dir, file_hint)
        sales = parse_bill_item_report(fp, branch_label)
        sales_data.extend(sales)
        print(f"  Bill Item Detailed Report [{branch_label}]: {len(sales)} item-day rows")

    for file_hint, branch_label in BRANCH_MAP.items():
        fp = _find_branch_file(discount_dir, file_hint)
        discounts = parse_discount_report(fp, branch_label, bill_totals_by_branch[branch_label])
        discount_data.extend(discounts)
        print(f"  Discount and Voucher Report [{branch_label}]: {len(discounts)} discount rows")

    dates = [b["date"] for b in all_bills if b["date"]]
    data_start = min(dates) if dates else None
    data_end = max(dates) if dates else None

    discount_reasons = sorted({d["reason"] for d in discount_data}) or \
        existing_data.get("meta", {}).get("discountReasons", [])

    old_meta = existing_data.get("meta", {})
    cancel_data = existing_data.get("cancelData", [])
    if cancel_data:
        print("  WARNING: no source report for item-level cancellations found in "
              "this export - keeping the previous cancelData unchanged "
              f"({len(cancel_data)} rows, stale).")
    else:
        print("  WARNING: no cancelData available (no prior data and no source "
              "report in this export). 'Cancelled KOTs' section will be empty.")

    return {
        "meta": {
            "dataStart": data_start,
            "dataEnd": data_end,
            "lastRefreshed": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
            "reasons": old_meta.get("reasons", ["Guest Cancellation", "Change Of Item", "Not Specified", "Other"]),
            "discountReasons": discount_reasons,
            "serviceTypes": ["Dine-In", "Home Delivery", "Takeaway", "Banquets & Catering"],
        },
        "allBills": all_bills,
        "salesData": sales_data,
        "discountData": discount_data,
        "cancelData": cancel_data,
    }


# --------------------------------------------------------------------------
# HTML injection
# --------------------------------------------------------------------------

DATA_LINE_RE = re.compile(r"^const DASHBOARD_DATA = .*;\s*$")


def load_existing_dashboard_data(html_path):
    if not os.path.exists(html_path):
        return {}
    with open(html_path, encoding="utf-8") as f:
        for line in f:
            if line.startswith("const DASHBOARD_DATA"):
                raw = line.strip()[len("const DASHBOARD_DATA = "):]
                if raw.endswith(";"):
                    raw = raw[:-1]
                return json.loads(raw)
    return {}


def write_dashboard_html(html_path, new_data):
    with open(html_path, encoding="utf-8") as f:
        lines = f.readlines()

    replaced = False
    new_line = "const DASHBOARD_DATA = " + json.dumps(new_data, ensure_ascii=False) + ";\n"
    for i, line in enumerate(lines):
        if DATA_LINE_RE.match(line):
            lines[i] = new_line
            replaced = True
            break
    if not replaced:
        raise RuntimeError("Could not find 'const DASHBOARD_DATA = ...;' line in the HTML file.")

    backup_path = html_path.replace(
        ".html", f".backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.html"
    )
    with open(html_path, encoding="utf-8") as f:
        original = f.read()
    with open(backup_path, "w", encoding="utf-8") as f:
        f.write(original)
    print(f"  Backed up previous HTML -> {backup_path}")

    with open(html_path, "w", encoding="utf-8") as f:
        f.writelines(lines)
    print(f"  Wrote updated dashboard -> {html_path}")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    posist_root, html_path = sys.argv[1], sys.argv[2]
    if not os.path.isdir(posist_root):
        print(f"Not a folder: {posist_root}")
        sys.exit(1)

    print(f"Reading raw reports from: {posist_root}")
    existing_data = load_existing_dashboard_data(html_path)
    new_data = build_dataset(posist_root, existing_data)

    print(f"\nTotals: {len(new_data['allBills'])} bills, "
          f"{len(new_data['salesData'])} item-day rows, "
          f"{len(new_data['discountData'])} discount rows, "
          f"{len(new_data['cancelData'])} cancellation rows "
          f"(period {new_data['meta']['dataStart']} to {new_data['meta']['dataEnd']}).\n")

    write_dashboard_html(html_path, new_data)
    print("\nDone.")


if __name__ == "__main__":
    main()
